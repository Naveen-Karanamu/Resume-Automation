import openpyxl
from collections import Counter
COLS = ["job_id","title","company","jd","keywords","prompt","response",
        "summary","ex_desc_1","ex_desc_2","resume_name","status","category"]
C = {n: i+1 for i, n in enumerate(COLS)}
ws = openpyxl.load_workbook("pipeline_data.xlsx").active
st = Counter(ws.cell(r, C["status"]).value for r in range(2, ws.max_row+1))
print("status counts:", dict(st))
print("\nNot-done rows:")
for r in range(2, ws.max_row+1):
    s = str(ws.cell(r, C["status"]).value or "")
    if s != "done":
        print(f"  row {r}: {ws.cell(r,C['job_id']).value} [{ws.cell(r,C['category']).value}] "
              f"({ws.cell(r,C['company']).value}) status={s} "
              f"kw={bool(ws.cell(r,C['keywords']).value)} resp={len(ws.cell(r,C['response']).value or '')}")
