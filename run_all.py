"""
ITERATIVE PIPELINE ORCHESTRATOR (resume-safe, Excel-backed)
===========================================================
Processes ALL jobs in the Apply Jobs table, one after another:

  Stage 0  Collect   : fetch jobs, iterate EVERY row, copy each JD, and record
                       job_id / title / company / jd into pipeline_data.xlsx
  Per job, sequentially:
    Stage 2  ATS      : paste JD -> FirstStep ATS Check -> copy Missing Keywords
    Stage 3  Sheet    : paste keywords(A9)+JD(B3) -> copy prompt(C3:C9)
    Stage 4  ChatGPT  : send prompt -> capture response
    Stage 5  Parse    : Summary / Resp1 / Resp2 -> append to resume_data.xlsx

ALL intermediate state lives in pipeline_data.xlsx (no .txt files). The run is
resume-safe: rows already 'done' are skipped, and a half-finished job resumes
from the first missing column. Per-job failures are logged and the run
continues to the next job.

pipeline_data.xlsx columns (header row 1):
  A job_id | B title | C company | D jd | E keywords | F prompt | G response
  H summary | I ex_desc_1 | J ex_desc_2 | K resume_name | L status

status: collected -> keywords -> prompt -> response -> done  (or error:<stage>)

Run:
    python run_all.py
"""

import json
import os
import re
import subprocess
import time
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

# ── CONFIG ──────────────────────────────────────────────────────────────────
APPLY_URL    = "https://firststepjob.com/application-incharge/apply-jobs?userId=google-oauth2|110590764416739071400"
BUILDER_URL  = "https://builder.firststepjob.com/resume-builder"
# SHEETS_URL is overridden at runtime from credentials.txt ('sheets_url=').
# This value is only a fallback if that line is absent.
SHEETS_URL   = "https://docs.google.com/spreadsheets/d/1BsXaJmFELvvB8s-bRogF9jNhiKgxvx-eUVcpBW-B6Ao/edit?gid=0#gid=0"
CHATGPT_URL  = "https://chatgpt.com/"
CHROME_EXE   = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
DEBUG_PORT   = 9222
HOME         = Path(os.path.expanduser("~"))
TEMP_DIR     = str(HOME / "AppData" / "Local" / "Temp" / "ChromeDebugSession")
HERE         = Path(__file__).parent
SESSION_FILE = HERE / "session.json"
CREDS_FILE   = HERE / "credentials.txt"
PIPELINE_XLSX = HERE / "pipeline_data.xlsx"
RESUME_XLSX   = HERE / "resume_data.xlsx"

CREATED_BY = "Job Collector"
KEYWORDS_CELL, JD_CELL, PROMPT_RANGE = "A9", "B3", "C3:C9"

# All filter inputs come ONLY from credentials.txt (no hardcoded values):
#   candidate -> "Select Candidate" dropdown + Resume Name
#   category  -> "Category" dropdown
#   date      -> "Date" filter (a FIXED value like 10/12/2026, NOT yesterday)

# Delay (seconds) between ChatGPT prompts to avoid rate limits.
CHATGPT_GAP = 5

# pipeline_data.xlsx column order (1-indexed). 'category' is last (col M) so
# older 12-column files still line up for columns A-L.
COLS = ["job_id", "title", "company", "jd", "keywords", "prompt",
        "response", "summary", "ex_desc_1", "ex_desc_2", "resume_name", "status",
        "category"]
C = {name: i + 1 for i, name in enumerate(COLS)}   # name -> column index
# ─────────────────────────────────────────────────────────────────────────────


# ╔═══════════════════ infrastructure ═══════════════════╗
def debug_port_up():
    try:
        urllib.request.urlopen(f"http://127.0.0.1:{DEBUG_PORT}/json", timeout=2)
        return True
    except Exception:
        return False


def launch_chrome():
    print("[*] Launching Chrome with remote debugging...")
    Path(TEMP_DIR).mkdir(parents=True, exist_ok=True)
    proc = subprocess.Popen([
        CHROME_EXE, f"--remote-debugging-port={DEBUG_PORT}",
        f"--user-data-dir={TEMP_DIR}",
        "--no-first-run", "--no-default-browser-check", "about:blank",
    ])
    for _ in range(30):
        if debug_port_up():
            print(f"[+] Chrome ready (PID {proc.pid}).")
            return proc
        time.sleep(1)
    raise RuntimeError("Chrome debug port did not come up.")


def read_credentials():
    """Return a dict of all key=value lines from credentials.txt."""
    creds = {}
    if CREDS_FILE.exists():
        for line in CREDS_FILE.read_text().splitlines():
            line = line.strip()
            if "=" in line and not line.startswith("#"):
                k, _, v = line.partition("=")
                creds[k.strip()] = v.strip()
    return creds


def load_firststep_cookies(context):
    if not SESSION_FILE.exists():
        return
    try:
        storage = json.loads(SESSION_FILE.read_text())
        cookies = [c for c in storage.get("cookies", [])
                   if "firststepjob.com" in c.get("domain", "")]
        if cookies:
            context.add_cookies(cookies)
            print(f"[+] Loaded {len(cookies)} FirstStep cookie(s).")
    except Exception as e:
        print(f"[!] cookie load failed: {e}")


def read_clipboard(page):
    try:
        return page.evaluate("navigator.clipboard.readText()")
    except Exception:
        try:
            return page.evaluate("""async () => {
                const ta=document.createElement('textarea');document.body.appendChild(ta);
                ta.focus();document.execCommand('paste');const v=ta.value;ta.remove();return v;}""")
        except Exception:
            return ""


# ╔═══════════════════ pipeline_data.xlsx store ═══════════════════╗
def pipe_load():
    """Load (or create) pipeline_data.xlsx and return (wb, ws)."""
    if PIPELINE_XLSX.exists():
        wb = openpyxl.load_workbook(PIPELINE_XLSX)
        ws = wb.active
        # Ensure every column header exists (adds the new 'category' header to
        # older 12-column files without disturbing existing data columns A-L).
        for i, name in enumerate(COLS, start=1):
            if ws.cell(1, i).value != name:
                ws.cell(1, i, name)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "pipeline"
        for i, name in enumerate(COLS, start=1):
            ws.cell(1, i, name)
    return wb, ws


def pipe_rows(ws):
    """Return list of (row_number, dict-of-values) for all data rows."""
    out = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, C["job_id"]).value in (None, ""):
            continue
        vals = {name: ws.cell(r, C[name]).value for name in COLS}
        out.append((r, vals))
    return out


def pipe_find_row(ws, job_id, category=None):
    """Find a row by job_id. When category is given, match (job_id, category)
    so the same job under different categories is treated as separate rows."""
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, C["job_id"]).value or "") != str(job_id):
            continue
        if category is None:
            return r
        if str(ws.cell(r, C["category"]).value or "").strip().lower() == \
           str(category).strip().lower():
            return r
    return None


def pipe_set(ws, row, **kwargs):
    for name, value in kwargs.items():
        ws.cell(row, C[name], value)


def pipe_save(wb):
    """Save with a retry so a brief Excel lock doesn't lose data."""
    for attempt in range(10):
        try:
            wb.save(PIPELINE_XLSX)
            return True
        except PermissionError:
            if attempt == 0:
                print("  [!] pipeline_data.xlsx is open/locked — waiting to save...")
            time.sleep(3)
    print("  [!] Could not save pipeline_data.xlsx (still locked).")
    return False


# ╔═══════════════════ Apply Jobs helpers ═══════════════════╗
def ensure_on_apply_page(page, settle=2.5):
    for _ in range(6):
        if "apply-jobs" not in page.url:
            page.goto(APPLY_URL, wait_until="domcontentloaded", timeout=60000)
        try:
            page.wait_for_load_state("networkidle", timeout=8000)
        except PlaywrightTimeoutError:
            pass
        try:
            page.wait_for_selector('[role="combobox"]', timeout=8000)
            time.sleep(settle)
            return True
        except PlaywrightTimeoutError:
            time.sleep(1)
    return False


def select_category(page, open_locator, target, label="Category"):
    """
    The Category field is a MULTI-select. Options already chosen stay selected,
    so before picking the target we must DESELECT every currently-selected
    option (click each one once), then select ONLY the target.

    MUI marks chosen options with aria-selected="true". We toggle those off,
    then toggle the target on, leaving exactly one category selected.
    """
    for attempt in range(2):
        try:
            ensure_on_apply_page(page, settle=0.5)
            open_locator().click()
            page.wait_for_selector('[role="listbox"], [role="option"]', timeout=6000)
            time.sleep(0.4)

            # 1) Deselect every currently-selected option whose text != target.
            # Re-query each pass because the DOM/selection updates after a click.
            for _ in range(10):
                sel = page.locator('[role="option"][aria-selected="true"]')
                cnt = sel.count()
                clicked_any = False
                for i in range(cnt):
                    opt = sel.nth(i)
                    try:
                        txt = (opt.inner_text() or "").strip()
                    except Exception:
                        continue
                    if txt.lower() != target.lower():
                        opt.click()
                        time.sleep(0.25)
                        print(f"  [-] {label}: deselected '{txt}'.")
                        clicked_any = True
                        break      # re-query after each deselect
                if not clicked_any:
                    break

            # 2) Select the target (exact text) if not already selected.
            tgt = page.locator('[role="option"]').filter(
                has_text=re.compile(rf'^{re.escape(target)}$', re.IGNORECASE)
            ).first
            if tgt.count() == 0:
                # fall back to a contains match if exact not found
                tgt = page.locator(f'[role="option"]:has-text("{target}")').first
            tgt.wait_for(state="visible", timeout=6000)
            is_sel = (tgt.get_attribute("aria-selected") or "").lower() == "true"
            if not is_sel:
                tgt.click()
                time.sleep(0.2)
            print(f"  [+] {label} = {target}.")

            # 3) Close the dropdown.
            page.keyboard.press("Escape")
            page.locator("body").click(position={"x": 10, "y": 10}, force=True)
            time.sleep(0.3)
            return True
        except Exception as e:
            print(f"  [~] {label} attempt {attempt+1}: {str(e)[:70]}")
            try:
                page.keyboard.press("Escape")
            except Exception:
                pass
            time.sleep(1)
    print(f"  [!] {label} could not be set to '{target}'.")
    return False


def pick_dropdown(page, open_locator, option_text, label=""):
    for attempt in range(2):
        try:
            ensure_on_apply_page(page, settle=0.5)
            open_locator().click()
            page.wait_for_selector('[role="listbox"], [role="option"]', timeout=6000)
            time.sleep(0.3)
            opt = page.locator(f'[role="option"]:has-text("{option_text}")').first
            opt.wait_for(state="visible", timeout=6000)
            opt.click()
            page.keyboard.press("Escape")
            page.locator("body").click(position={"x": 10, "y": 10}, force=True)
            time.sleep(0.3)
            print(f"  [+] {label} = {option_text}.")
            return True
        except Exception as e:
            print(f"  [~] {label} attempt {attempt+1}: {str(e)[:70]}")
            try:
                page.keyboard.press("Escape")
            except Exception:
                pass
            time.sleep(1)
    return False


def setup_filters(page, category=None, filter_date=None, candidate=None):
    """Fill the filters in order: Created By -> Date -> Category -> Candidate,
    then click Fetch Jobs.

    date/category/candidate come from credentials.txt. The date is used EXACTLY
    as given (e.g. '10/12/2026') — it is no longer auto-set to yesterday. If no
    date is provided, fall back to yesterday for safety.
    """
    category = category or ""
    candidate = candidate or ""

    # Resolve the date string to type into the picker.
    if filter_date:
        date_str = filter_date.strip()
    else:
        y = datetime.today() - timedelta(days=1)
        date_str = y.strftime("%m/%d/%Y")

    # Parse MM/DD/YYYY so we can type the digits into the MUI date picker.
    m = re.match(r'\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{2,4})\s*', date_str)
    if m:
        mm = m.group(1).zfill(2)
        dd = m.group(2).zfill(2)
        yyyy = m.group(3)
        if len(yyyy) == 2:
            yyyy = "20" + yyyy
    else:
        print(f"  [!] Date '{date_str}' not MM/DD/YYYY; using yesterday.")
        y = datetime.today() - timedelta(days=1)
        mm, dd, yyyy = y.strftime("%m"), y.strftime("%d"), y.strftime("%Y")
        date_str = y.strftime("%m/%d/%Y")

    print(f"[*] Filters -> Created By='{CREATED_BY}', Date='{date_str}', "
          f"Category='{category}', Candidate='{candidate}'")

    # Created By
    pick_dropdown(page, lambda: page.locator('[role="combobox"]').nth(0), CREATED_BY, "Created By")

    # Date (MUI date picker: type MM then DD then YYYY)
    try:
        ensure_on_apply_page(page, settle=0.5)
        page.locator('[role="spinbutton"][aria-label="Month"]').first.click()
        time.sleep(0.2)
        for part in (mm, dd, yyyy):
            page.keyboard.type(part); time.sleep(0.2)
        page.keyboard.press("Escape")
        print(f"  [+] Date = {mm}/{dd}/{yyyy}.")
    except Exception as e:
        print(f"  [!] Date failed: {str(e)[:70]}")

    # Category (MULTI-select: deselect any prior choices, then select target)
    select_category(page, lambda: page.locator(
        '.MuiFormControl-root:has(label:has-text("Category")) [role="combobox"]').first,
        category, "Category")

    # Candidate ("Select Candidate" dropdown)
    try:
        cand_combo = page.locator(
            '.MuiFormControl-root:has(label:has-text("Candidate")) [role="combobox"]'
        ).first
        if cand_combo.count() > 0:
            pick_dropdown(page, lambda: cand_combo, candidate, "Candidate")
        else:
            print("  [~] No Candidate dropdown found (may be pre-selected).")
    except Exception as e:
        print(f"  [!] Candidate select failed: {str(e)[:70]}")

    # Fetch Jobs
    try:
        btn = page.locator('button:has-text("Fetch Jobs")').first
        for _ in range(20):
            if btn.is_enabled():
                break
            time.sleep(0.5)
        btn.scroll_into_view_if_needed(); btn.click()
        print("  [+] Fetch Jobs clicked.")
        time.sleep(4)
    except Exception as e:
        print(f"  [!] Fetch Jobs failed: {e}")
    page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    time.sleep(1)


def parse_yoe(text):
    """Extract the numeric years from a YOE cell like '3 yrs', '6 yrs', '2 yr'.
    Returns an int, or None if no number is present (e.g. '-')."""
    if not text:
        return None
    m = re.search(r'\d+', text)
    return int(m.group()) if m else None


def read_row_meta(page, row_index):
    """Read job_id (td0), title (td1), company (td2), YOE (td4) for a row."""
    rows = page.locator('table tbody tr').all()
    if row_index >= len(rows):
        return None
    cells = rows[row_index].locator('td').all()
    def cell(i):
        return " ".join((cells[i].inner_text() or "").split()).strip() if i < len(cells) else ""
    yoe_text = cell(4)   # YOE column, e.g. "3 yrs"
    return {
        "job_id": cell(0),
        "title": cell(1).rstrip(","),
        "company": cell(2),
        "yoe_text": yoe_text,
        "yoe": parse_yoe(yoe_text),
        "n_rows": len(rows),
    }


def copy_jd_for_row(page, row_index):
    """Open the Resume modal for the given row, click the copy icon, return JD."""
    ensure_on_apply_page(page, settle=0.5)
    rows = page.locator('table tbody tr').all()
    if row_index >= len(rows):
        return ""
    row = rows[row_index]
    page.evaluate("document.querySelector('table')?.parentElement?.scrollBy(600,0)")
    time.sleep(0.3)
    rbtn = row.locator('button:text-is("Resume")').first
    if rbtn.count() == 0:
        for b in row.locator('button').all():
            if b.inner_text().strip() == "Resume":
                rbtn = b; break
    try:
        rbtn.scroll_into_view_if_needed()
        time.sleep(0.3)
        rbtn.evaluate("el => el.click()")
    except Exception:
        try:
            rbtn.click(force=True)
        except Exception as e:
            print(f"    [!] Resume click failed: {str(e)[:80]}")
            return ""
    time.sleep(2)
    try:
        page.wait_for_selector('[role="dialog"]', timeout=8000)
    except PlaywrightTimeoutError:
        pass
    time.sleep(0.5)
    jd = ""
    try:
        b = page.locator('[aria-label="Copy job description"], [title="Copy job description"], '
                         '[aria-label*="Copy job description" i]').first
        b.wait_for(state="visible", timeout=8000)
        b.scroll_into_view_if_needed(); b.click()
        time.sleep(0.6)
        jd = read_clipboard(page)
    except Exception as e:
        print(f"    [!] copy JD failed: {str(e)[:90]}")
    # Close the modal so the next row's Resume button is clickable
    try:
        page.keyboard.press("Escape")
        time.sleep(0.5)
    except Exception:
        pass
    return jd or ""


# ╔═══════════════════ Builder ATS ═══════════════════╗
def builder_login_if_needed(page, email, password):
    page.goto(BUILDER_URL, wait_until="domcontentloaded", timeout=60000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except PlaywrightTimeoutError:
        pass
    time.sleep(2)
    ef = page.locator('#email, input[type="email"], input[name="email"]').first
    pf = page.locator('#password, input[type="password"], input[name="password"]').first
    if ef.count() and pf.count() and ef.is_visible():
        print("  [*] Logging into builder...")
        ef.fill(email); pf.fill(password)
        page.locator('button[type="submit"], button:has-text("Log in"), '
                     'button:has-text("Login"), button:has-text("Sign in"), '
                     'button:has-text("Continue")').first.click()
        try:
            page.wait_for_load_state("networkidle", timeout=25000)
        except PlaywrightTimeoutError:
            pass
        time.sleep(4)
    else:
        print("  [+] Builder already logged in.")


def ats_keywords(page, jd):
    """Open ATS Check (if needed), paste JD, run, copy Missing Keywords.
    Clicks 'Analyze Another Job' first if a previous result is showing so the
    paste box is available for this job."""
    # If a previous analysis result is on screen, reset it.
    try:
        again = page.locator('button:has-text("Analyze Another Job")').first
        if again.count() > 0 and again.is_visible():
            again.click()
            time.sleep(1)
    except Exception:
        pass

    already_open = page.locator(
        'textarea[placeholder*="Paste the complete job description" i]'
    ).count() > 0
    if not already_open:
        try:
            b = page.locator('button:text-is("ATS Check")').first
            if b.count() == 0:
                b = page.locator('button:has-text("ATS Check")').first
            b.wait_for(state="visible", timeout=25000)
            b.scroll_into_view_if_needed(); b.click()
            print("    [+] ATS Check opened.")
        except Exception as e:
            print(f"    [!] ATS Check click failed: {str(e)[:90]}"); return ""
    try:
        page.wait_for_selector('[role="dialog"], .MuiDialog-root', timeout=15000)
    except PlaywrightTimeoutError:
        pass
    time.sleep(1)
    try:
        ta = page.locator('textarea[placeholder*="Paste the complete job description" i], '
                          'textarea[placeholder*="job description" i]').first
        if ta.count() == 0:
            ta = page.locator('[role="dialog"] textarea, .MuiDialog-root textarea').first
        ta.wait_for(state="visible", timeout=10000)
        ta.click(); ta.fill(jd)
        print(f"    [+] Pasted JD ({len(jd)} chars).")
        time.sleep(0.5)
    except Exception as e:
        print(f"    [!] paste JD failed: {str(e)[:90]}"); return ""
    try:
        r = page.locator('button:has-text("FirstStep ATS Check")').first
        r.wait_for(state="visible", timeout=10000); r.click()
        print("    [+] FirstStep ATS Check running...")
    except Exception as e:
        print(f"    [!] run failed: {str(e)[:90]}"); return ""
    try:
        page.wait_for_selector('text=/Missing Keywords/i', timeout=90000)
        time.sleep(1)
    except PlaywrightTimeoutError:
        print("    [!] Missing Keywords not found in 90s.")
    kw = ""
    try:
        page.locator('text=/Missing Keywords/i').first.scroll_into_view_if_needed()
        time.sleep(0.5)
        cb = page.locator('[aria-label*="Copy missing" i], [title*="Copy missing" i], '
                         '[aria-label*="missing keyword" i], [title*="missing keyword" i]').first
        if cb.count() == 0:
            cb = page.locator('xpath=//*[contains(translate(text(),"MISSING","missing"),"missing keywords")]'
                             '/following::button[1]').first
        cb.wait_for(state="visible", timeout=10000)
        cb.scroll_into_view_if_needed(); cb.click()
        time.sleep(0.6)
        kw = read_clipboard(page)
    except Exception as e:
        print(f"    [!] copy keywords failed: {str(e)[:90]}")
    return kw or ""


# ╔═══════════════════ Google Sheet ═══════════════════╗
def set_clipboard(page, text):
    page.evaluate("async (t)=>{await navigator.clipboard.writeText(t);}", text)


def go_to_cell(page, ref):
    nb = page.locator('input.waffle-name-box, input[aria-label="Name box"], '
                     'input[aria-label="Name Box"]').first
    nb.click(); nb.press("Control+a"); nb.type(ref)
    page.keyboard.press("Enter"); time.sleep(0.5)


def clear_cell(page, ref):
    go_to_cell(page, ref)
    page.keyboard.press("Delete"); time.sleep(0.3)


def paste_cell(page, ref, value):
    set_clipboard(page, value)
    go_to_cell(page, ref)
    page.keyboard.press("Delete"); time.sleep(0.2)
    page.keyboard.press("F2"); time.sleep(0.3)
    page.keyboard.press("Control+v"); time.sleep(0.6)
    page.keyboard.press("Enter"); time.sleep(0.5)


def get_sheet_page(context):
    page = None
    for pg in context.pages:
        if "docs.google.com/spreadsheets" in pg.url and "/edit" in pg.url:
            page = pg; break
    if page is None:
        page = context.new_page()
        page.goto(SHEETS_URL, wait_until="domcontentloaded", timeout=60000)
    for _ in range(20):
        try:
            page.wait_for_load_state("networkidle", timeout=5000)
        except PlaywrightTimeoutError:
            pass
        if "/edit" in page.url and "accounts.google.com" not in page.url:
            break
        time.sleep(1)
    page.bring_to_front(); time.sleep(1.5)
    if "accounts.google.com" in page.url:
        return None
    try:
        page.wait_for_selector('input.waffle-name-box, [aria-label="Name box"], [aria-label="Name Box"]', timeout=20000)
    except PlaywrightTimeoutError:
        return None
    return page


def _norm_tab(s):
    """Normalise a tab/category name for matching: lowercase, strip, remove
    all non-alphanumeric chars (so 'AI ML Engineer' == 'AIML ENGINEER')."""
    return re.sub(r'[^a-z0-9]', '', (s or "").lower())


def select_sheet_tab(page, category):
    """Click the bottom sheet tab whose name matches `category`.
    Returns True if a matching tab was found and activated, else False.
    Matching is case/space-insensitive ('AI ML Engineer' -> 'AIML ENGINEER')."""
    target = _norm_tab(category)
    if not target:
        return False
    try:
        # Google Sheets bottom tabs: each sheet tab name sits in an element with
        # class 'docs-sheet-tab-name' inside a 'docs-sheet-tab' container.
        tabs = page.locator('.docs-sheet-tab-name').all()
        names = []
        for t in tabs:
            try:
                nm = (t.inner_text() or "").strip()
            except Exception:
                nm = ""
            names.append(nm)
            if _norm_tab(nm) == target:
                t.scroll_into_view_if_needed()
                t.click()
                time.sleep(1.2)
                print(f"    [+] Switched to sheet tab '{nm}'.")
                return True
        print(f"    [!] No sheet tab matches category '{category}'. "
              f"Available tabs: {names}")
        return False
    except Exception as e:
        print(f"    [!] select_sheet_tab failed: {str(e)[:90]}")
        return False


def sheet_get_prompt(page, keywords, jd, category=None):
    # Switch to the sheet tab for this category before pasting / reading.
    if category:
        select_sheet_tab(page, category)
    clear_cell(page, KEYWORDS_CELL)
    paste_cell(page, KEYWORDS_CELL, keywords)
    paste_cell(page, JD_CELL, jd)
    time.sleep(1.5)
    go_to_cell(page, PROMPT_RANGE)
    page.keyboard.press("Control+c"); time.sleep(0.8)
    prompt = (read_clipboard(page) or "").strip()

    # Clean up the sheet's input cells so the next job/category starts fresh
    # (keywords in A9 and JD in B3 don't carry over).
    try:
        clear_cell(page, KEYWORDS_CELL)
        clear_cell(page, JD_CELL)
        print("    [+] Cleared sheet input cells (A9, B3).")
    except Exception as e:
        print(f"    [~] Sheet cleanup failed: {str(e)[:70]}")

    return prompt


# ╔═══════════════════ ChatGPT ═══════════════════╗
def _extract_chatgpt_answer(page):
    """Return the best answer block from ChatGPT, PRESERVING bold words as
    markdown '**word**'. ChatGPT renders bold as <strong>/<b> elements; plain
    innerText would drop that styling, so we walk the DOM and re-wrap bold runs
    in ** markers. Line structure is kept (one line per block element)."""
    return page.evaluate(r"""() => {
        const want=/professional summary|responsibilit|experience/i;
        const isPrompt = t => /STRICT INSTRUCTIONS|Forget any previous job/i.test(t);
        let container=null, bestLen=0;
        document.querySelectorAll('div,section,article').forEach(el=>{
            const t=el.innerText||'';
            if(t.length>200 && t.length<30000 && want.test(t) && !isPrompt(t)){
                if(t.length>bestLen){ bestLen=t.length; container=el; }
            }
        });
        if(!container) return '';

        // Recursively serialise a node to text, wrapping bold runs in **..**.
        function serialize(node){
            if(node.nodeType === Node.TEXT_NODE) return node.nodeValue;
            if(node.nodeType !== Node.ELEMENT_NODE) return '';
            const tag = node.tagName.toLowerCase();
            // gather children text
            let inner = '';
            node.childNodes.forEach(ch => { inner += serialize(ch); });
            // Only treat explicit <strong>/<b> tags as bold — that is exactly
            // what markdown '**word**' renders to. Do NOT use computed
            // font-weight: ChatGPT's UI chrome and section headings are heavy
            // and would be falsely bolded. Also avoid double-wrapping when a
            // bold tag is nested inside another bold tag.
            const isBold = (tag === 'strong' || tag === 'b');
            const alreadyBold = node.closest &&
                node.parentElement && node.parentElement.closest &&
                node.parentElement.closest('strong,b');
            if(isBold && !alreadyBold){
                const trimmed = inner.replace(/\*\*/g, '').trim();
                if(!trimmed) return inner;
                const lead = inner.match(/^\s*/)[0];
                const tail = inner.match(/\s*$/)[0];
                return lead + '**' + trimmed + '**' + tail;
            }
            return inner;
        }

        // Build text block-by-block so each heading/paragraph/li is its own
        // line. Include leaf <div>s too (the summary paragraph often sits in a
        // bare div, not a <p>) so we don't lose it.
        const parts = [];
        const seen = new Set();
        const blocks = container.querySelectorAll('h1,h2,h3,h4,h5,h6,p,li,div');
        if(blocks.length === 0){
            return serialize(container);
        }
        blocks.forEach(b => {
            const tag = b.tagName.toLowerCase();
            if(tag === 'p' && b.closest('li')) return;
            if(tag === 'div' && b.querySelector('p,li,div,h1,h2,h3,h4,h5,h6')) return; // not a leaf
            let line = serialize(b).replace(/[ \t]+/g,' ').trim();
            // Normalise bold markers: collapse accidental doubles and empties.
            line = line.replace(/\*\*\*\*/g, '')      // **** -> (nothing)
                       .replace(/\*\*\s*\*\*/g, ' ')  // ** ** -> space
                       .replace(/\*\*\*/g, '**');     // *** -> **
            if(line && !seen.has(line)){ seen.add(line); parts.push(line); }
        });
        return parts.join('\n');
    }""")


def chatgpt_answer(context, prompt):
    """Open a FRESH ChatGPT chat for this prompt, send it, and wait for a NEW
    response (different from whatever was previously on screen) to stabilise.
    Starting a new chat per job prevents returning a stale/previous answer."""
    if prompt.startswith('"') and prompt.endswith('"'):
        prompt = prompt[1:-1].strip()

    page = None
    for pg in context.pages:
        if "chatgpt.com" in pg.url or "chat.openai.com" in pg.url:
            page = pg; break
    if page is None:
        page = context.new_page()
    page.bring_to_front()

    # Always start a brand-new conversation so there is no previous answer to
    # mistakenly read back. chatgpt.com/ (root) opens a fresh chat.
    try:
        page.goto(CHATGPT_URL, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_load_state("networkidle", timeout=15000)
    except PlaywrightTimeoutError:
        pass
    time.sleep(2)
    if "auth" in page.url.lower() or "login" in page.url.lower():
        print("    [!] ChatGPT needs login."); return ""

    # Baseline: whatever answer-like text exists right now (should be empty on a
    # fresh chat, but guard anyway).
    baseline = (_extract_chatgpt_answer(page) or "").strip()

    try:
        box = page.locator('#prompt-textarea, div[contenteditable="true"], textarea').first
        box.wait_for(state="visible", timeout=20000)
        box.click()
        try:
            box.fill(prompt)
        except Exception:
            page.keyboard.insert_text(prompt)
        time.sleep(0.5)
        page.keyboard.press("Enter")
        print("    [+] Prompt sent (fresh chat).")
    except Exception as e:
        print(f"    [!] could not send prompt: {str(e)[:90]}"); return ""

    print("    [*] Waiting for NEW response (stabilising)...")
    last_len, stable, answer = -1, 0, ""
    for _ in range(180):          # up to ~6 min
        time.sleep(2)
        current = (_extract_chatgpt_answer(page) or "").strip()
        # Ignore until the answer is genuinely new vs the baseline.
        if not current or current == baseline:
            continue
        answer = current
        if len(answer) == last_len and len(answer) > 0:
            stable += 1
            if stable >= 3:        # unchanged for ~6s -> finished streaming
                break
        else:
            stable = 0
        last_len = len(answer)
    print(f"    [+] Response captured ({len(answer)} chars).")
    return answer


# ╔═══════════════════ keyword bolding ═══════════════════╗
def _split_keywords(kw_text):
    """Split the keywords cell into a cleaned, de-duplicated, longest-first list."""
    if not kw_text:
        return []
    seen, kws = set(), []
    for k in str(kw_text).split(","):
        k = k.strip()
        if len(k) < 2:
            continue
        low = k.lower()
        if low in seen:
            continue
        seen.add(low)
        kws.append(k)
    kws.sort(key=len, reverse=True)   # phrases before their component words
    return kws


def bold_keywords(response, kw_text):
    """Wrap each ATS keyword that appears as PLAIN text in `response` with
    markdown **bold**. Skips text already inside **...** (no double-bolding),
    matches whole words/phrases case-insensitively, and is idempotent."""
    keywords = _split_keywords(kw_text)
    if not response or not keywords:
        return response

    bold_pat = re.compile(r'\*\*.+?\*\*', re.DOTALL)

    def bold_one(text, kw):
        esc = re.escape(kw)
        left = r'(?<!\w)' if kw[:1].isalnum() else ''
        right = r'(?!\w)' if kw[-1:].isalnum() else ''
        return re.compile(left + esc + right, re.IGNORECASE).sub(
            lambda mm: "**" + mm.group(0) + "**", text)

    # split into alternating plain / already-bold segments; only bold plain
    segments, pos = [], 0
    for m in bold_pat.finditer(response):
        if m.start() > pos:
            segments.append(("plain", response[pos:m.start()]))
        segments.append(("bold", m.group(0)))
        pos = m.end()
    if pos < len(response):
        segments.append(("plain", response[pos:]))

    out = []
    for kind, text in segments:
        if kind == "bold":
            out.append(text); continue
        seg = text
        for kw in keywords:
            # re-protect spans newly bolded within this segment
            parts, p = [], 0
            for mm in bold_pat.finditer(seg):
                if mm.start() > p:
                    parts.append(("plain", seg[p:mm.start()]))
                parts.append(("bold", mm.group(0)))
                p = mm.end()
            if p < len(seg):
                parts.append(("plain", seg[p:]))
            seg = "".join(t if k2 == "bold" else bold_one(t, kw)
                          for k2, t in parts)
        out.append(seg)
    return "".join(out)


# ╔═══════════════════ parse + resume_data.xlsx ═══════════════════╗
def parse_sections(text):
    TRAILING = re.compile(
        r'^(Job Description Skills Coverage|Additional ATS Skills|Important accuracy note|'
        r'Important ATS|Important notes?\b|Before submitting|ATS alignment|'
        r'Skills Coverage|Job requirement|Accuracy\s*:|Compliance\s*:|'
        r'Technical keywords?\s*:|Additional requirement\s*:|Note\s*:|Notes?\b|'
        r'The following job-description)',
        re.IGNORECASE)

    # Strip a leading "ChatGPT said:" prefix (it can appear on its own line OR
    # inline at the start of the summary line, e.g. "ChatGPT said: Data ...").
    raw_lines = []
    for ln in text.splitlines():
        ln = ln.strip()
        ln = re.sub(r'^chatgpt said\s*:?\s*', '', ln, flags=re.IGNORECASE)
        raw_lines.append(ln)
    # drop noise lines: empties and 'Writing' artifacts
    lines = [ln for ln in raw_lines if ln and ln.lower() != "writing"]

    # A role header marks the start of an experience block. Observed formats:
    #   "Responsibilities 1: Techzert AI Engineer | New York, NY | Nov 2024 - Present"
    #   "Experience 1: ..."
    #   "Techzert — AI Engineer | New York, NY"   (date on next line)
    #   "Techzert — AI Engineer"                   (date on next line)
    #   "Techzert | AI Engineer"
    def is_date_line(ln):
        # e.g. "Nov 2024 - Present", "Jun 2022 - Jul 2023"
        return bool(re.match(
            r'^[A-Za-z]{3,9}\.?\s*\d{4}\s*[–\-]\s*(present|[A-Za-z]{3,9}\.?\s*\d{4})',
            ln, re.IGNORECASE))

    def is_summary_heading(ln):
        # A summary heading can be "Professional Summary", "Summary", or a
        # category-prefixed form like "AIML ENGINEER — PROFESSIONAL SUMMARY".
        # It's a short heading (not a sentence) that contains 'summary'.
        if ln.rstrip().endswith((".", "%", ")")):
            return False
        return bool(re.search(r'\bprofessional\s+summary\b', ln, re.IGNORECASE)
                    or re.match(r'^summary\s*$', ln, re.IGNORECASE))

    def is_role_header(ln):
        # Never treat a standalone date range as a role header (its dash would
        # otherwise look like a "Company – Title" separator).
        if is_date_line(ln):
            return False
        # Never treat a summary heading as a role header, even if it contains a
        # dash (e.g. "AIML ENGINEER — PROFESSIONAL SUMMARY").
        if is_summary_heading(ln):
            return False
        # (a) explicit "Responsibilities N:" / "Experience N:" label
        if re.match(r'^(responsibilit(?:y|ies)|experience)\s*\d*\s*[:\-]',
                    ln, re.IGNORECASE):
            return True
        # (b) "Company <dash> Title" heading style: has an em/en-dash separator,
        #     is not a full sentence, and is short once any trailing
        #     "| city | date" tail is removed.
        head = re.split(r'\s\|\s', ln)[0]   # text before the first " | "
        if re.search(r'.+\s[—–]\s.+', head) and len(head) < 60 \
                and not ln.rstrip().endswith((".", "%", ")")):
            return True
        return False

    def is_section_label(ln):
        # standalone section wrappers that are NOT role headers
        return bool(re.match(r'^(professional\s+experience|work\s+experience)\s*$',
                             ln, re.IGNORECASE))

    def is_meta(ln):
        # pure meta lines to skip inside bullet collection (city/date-only,
        # 'Rewritten...' notes, section wrappers). Role headers are handled
        # separately and must be checked BEFORE this.
        if re.match(r'^(Rewritten|Revised|Reworded)\b', ln, re.IGNORECASE):
            return True
        # a standalone date line like "Nov 2024 - Present" or "Jun 2022 - Jul 2023"
        if re.match(r'^[A-Za-z]{3,9}\.?\s*\d{4}\s*[–\-]\s*(present|[A-Za-z]{3,9}\.?\s*\d{4})',
                    ln, re.IGNORECASE):
            return True
        if is_section_label(ln):
            return True
        return False

    def is_bullet(ln):
        return ln.rstrip().endswith((".", "%", ")")) and len(ln) > 25

    role_idxs = []
    for i, ln in enumerate(lines):
        if TRAILING.match(ln):
            break
        if is_summary_heading(ln):
            continue
        # Role headers take precedence over meta classification.
        if is_role_header(ln):
            role_idxs.append(i)

    def clean_bullet(ln):
        # Strip a leading bullet glyph/numbering, but DO NOT strip markdown
        # bold markers (**). Only remove a single leading •, -, ●, ▪, ◦ or
        # "1." style numbering — never a '*' that is part of '**bold**'.
        ln = re.sub(r'^\s*[\u2022\u25cf\u25aa\u25e6]\s*', '', ln)
        ln = re.sub(r'^\s*-\s+', '', ln)          # leading dash bullet (with space)
        ln = re.sub(r'^\s*\d+[\.\)]\s*', '', ln)
        return ln.strip()

    def bullets_after(idx):
        out = []
        for ln in lines[idx + 1:]:
            if TRAILING.match(ln):
                break
            # Stop at the next role header (start of the other experience).
            if is_role_header(ln):
                break
            if is_meta(ln):
                continue
            if is_bullet(ln):
                out.append(clean_bullet(ln))
        return "\n".join(out).strip()

    resp1 = bullets_after(role_idxs[0]) if len(role_idxs) >= 1 else ""
    resp2 = bullets_after(role_idxs[1]) if len(role_idxs) >= 2 else ""

    summary_lines = []
    end = role_idxs[0] if role_idxs else len(lines)
    for ln in lines[:end]:
        if is_summary_heading(ln):
            continue
        if is_meta(ln):
            continue
        summary_lines.append(ln)
    summary = " ".join(summary_lines).strip()
    return summary, resp1, resp2


_BOLD_RE = re.compile(r'\*\*(.+?)\*\*', re.DOTALL)


def to_rich_text(text):
    """Convert a string containing markdown **bold** markers into an openpyxl
    CellRichText so the bold words render bold inside the Excel cell. If there
    are no ** markers, return the plain string unchanged."""
    if not text or "**" not in text:
        return text
    parts = []
    pos = 0
    bold_font = InlineFont(b=True)
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            parts.append(text[pos:m.start()])        # normal text
        parts.append(TextBlock(bold_font, m.group(1)))  # bold run
        pos = m.end()
    if pos < len(text):
        parts.append(text[pos:])
    if not parts:
        return text
    return CellRichText(*parts)


def append_resume_data(summary, resp1, resp2, resume_name="", job_title=""):
    """Append one row to resume_data.xlsx with wrap-text and tall row height.

    IMPORTANT: the Summary / Ex Desc cells keep the markdown **bold** markers as
    PLAIN TEXT (they are NOT converted to Excel rich-text). resume_automation.py
    reads these markers and recreates REAL bold inside the resume builder, which
    is what makes the keywords bold in the final PDF. (Excel will show literal
    ** around keywords, but this file is just the intermediate hand-off.)"""
    wb = openpyxl.load_workbook(RESUME_XLSX)
    ws = wb.active
    ws.append([job_title, "", summary, resp1, resp2, resume_name])
    row = ws.max_row
    wrap = Alignment(wrap_text=True, vertical="top")
    for col in (3, 4, 5):
        ws.cell(row=row, column=col).alignment = wrap
    for cl, w in (("C", 60), ("D", 70), ("E", 70)):
        ws.column_dimensions[cl].width = w
    max_lines = max(resp1.count("\n") + 1, resp2.count("\n") + 1)
    ws.row_dimensions[row].height = max(90, max_lines * 45)
    for attempt in range(10):
        try:
            wb.save(RESUME_XLSX)
            return row
        except PermissionError:
            if attempt == 0:
                print("    [!] resume_data.xlsx locked — waiting to save...")
            time.sleep(3)
    print("    [!] Could not save resume_data.xlsx (still locked).")
    return None


# ╔═══════════════════ STAGE 0: collect all jobs ═══════════════════╗
def collect_all_jobs(apply_page, wb, ws, category=None, filter_date=None,
                     candidate=None, max_experience=None):
    """Fetch jobs and record every row's job_id/title/company/jd into
    pipeline_data.xlsx. Rows already present (by job_id) are not re-collected.
    If max_experience is set, jobs whose YOE exceeds it are skipped."""
    if not ensure_on_apply_page(apply_page, settle=2.5):
        print("[!] Apply Jobs not ready."); return 0
    setup_filters(apply_page, category=category, filter_date=filter_date, candidate=candidate)
    ensure_on_apply_page(apply_page, settle=1.0)

    rows = apply_page.locator('table tbody tr').all()
    n = len(rows)
    cap = f", YOE <= {max_experience}" if max_experience is not None else ""
    print(f"[*] [{category}] Found {n} job(s){cap}. Collecting JDs...")
    collected = 0
    skipped_yoe = 0
    for i in range(n):
        meta = read_row_meta(apply_page, i)
        if not meta or not meta["job_id"]:
            print(f"  [~] Row {i+1}: no job_id, skipping.")
            continue
        jid = meta["job_id"]

        # YOE filter: skip jobs requiring more experience than the candidate has.
        # Jobs with no numeric YOE (e.g. '-') are kept (not filtered out).
        if max_experience is not None and meta["yoe"] is not None \
                and meta["yoe"] > max_experience:
            print(f"  [x] {i+1}/{n} {jid} YOE={meta['yoe_text']} "
                  f"> {max_experience} — skip.")
            skipped_yoe += 1
            continue

        # category-aware dedup: same job under a different category is separate
        existing = pipe_find_row(ws, jid, category=category)
        if existing and ws.cell(existing, C["jd"]).value:
            print(f"  [=] {i+1}/{n} {jid} [{category}] already collected — skip.")
            continue
        print(f"  [>] {i+1}/{n} {jid} | {meta['title']} | {meta['company']} "
              f"| YOE={meta['yoe_text']} [{category}]")
        jd = copy_jd_for_row(apply_page, i)
        if not jd:
            print(f"      [!] JD empty for {jid}; recording meta only.")
        row = existing or (ws.max_row + 1)
        pipe_set(ws, row,
                 job_id=jid, title=meta["title"], company=meta["company"],
                 jd=jd, category=category,
                 status="collected" if jd else "error:collect")
        pipe_save(wb)
        collected += 1
    print(f"[+] [{category}] Collection done: {collected} new row(s), "
          f"{skipped_yoe} skipped by YOE cap.")
    return collected


# ╔═══════════════════ per-job processing ═══════════════════╗
def process_job(context, builder_page, sheet_page, wb, ws, row, candidate=None):
    """Run stages 2-5 for one pipeline row, resuming from where it left off.
    Writes results back to pipeline_data.xlsx after each stage. The resume Job
    Title comes from this row's stored category."""
    candidate = candidate or ""
    category = ws.cell(row, C["category"]).value or ""
    jid     = ws.cell(row, C["job_id"]).value
    title   = ws.cell(row, C["title"]).value or ""
    company = ws.cell(row, C["company"]).value or ""
    jd      = ws.cell(row, C["jd"]).value or ""
    if not jd:
        print(f"  [!] {jid}: no JD stored — cannot process.")
        pipe_set(ws, row, status="error:no_jd"); pipe_save(wb); return False

    # Stage 2: keywords
    keywords = ws.cell(row, C["keywords"]).value or ""
    if not keywords:
        print(f"  [2] {jid}: ATS keywords...")
        keywords = ats_keywords(builder_page, jd)
        if not keywords:
            pipe_set(ws, row, status="error:keywords"); pipe_save(wb); return False
        pipe_set(ws, row, keywords=keywords, status="keywords"); pipe_save(wb)

    # Stage 3: prompt
    prompt = ws.cell(row, C["prompt"]).value or ""
    if not prompt:
        print(f"  [3] {jid}: Sheet prompt...")
        if sheet_page is None:
            pipe_set(ws, row, status="error:sheet_login"); pipe_save(wb); return False
        prompt = sheet_get_prompt(sheet_page, keywords, jd, category=category)
        if not prompt:
            pipe_set(ws, row, status="error:prompt"); pipe_save(wb); return False
        pipe_set(ws, row, prompt=prompt, status="prompt"); pipe_save(wb)

    # Stage 4: response
    response = ws.cell(row, C["response"]).value or ""
    if not response:
        print(f"  [4] {jid}: ChatGPT...")
        response = chatgpt_answer(context, prompt)
        if not response:
            pipe_set(ws, row, status="error:response"); pipe_save(wb); return False
        pipe_set(ws, row, response=response, status="response"); pipe_save(wb)
        time.sleep(CHATGPT_GAP)

    # Stage 5: auto-bold ATS keywords, then parse + append to resume_data.xlsx
    print(f"  [5] {jid}: parse & write resume_data...")
    bolded = bold_keywords(response, keywords)
    if bolded != response:
        response = bolded
        pipe_set(ws, row, response=response); pipe_save(wb)
    summary, resp1, resp2 = parse_sections(response)
    if not (summary and resp1 and resp2):
        print(f"      [!] parse incomplete (S{len(summary)} R1{bool(resp1)} R2{bool(resp2)}).")
        pipe_set(ws, row, summary=summary, ex_desc_1=resp1, ex_desc_2=resp2,
                 status="error:parse"); pipe_save(wb)
        return False
    resume_name = f"{candidate} - Resume {company}".strip() if company \
                  else f"{candidate} - Resume"
    # Job Title (resume_data column A) = the CATEGORY from credentials.txt,
    # NOT the per-job title scraped from the jobs screen.
    job_title = category
    append_resume_data(summary, resp1, resp2, resume_name, job_title)
    pipe_set(ws, row, summary=summary, ex_desc_1=resp1, ex_desc_2=resp2,
             resume_name=resume_name, status="done"); pipe_save(wb)
    print(f"      [+] DONE {jid}: title='{job_title}' -> {resume_name} "
          f"(R1 {resp1.count(chr(10))+1} lines, R2 {resp2.count(chr(10))+1} lines).")
    return True


# ╔═══════════════════ MAIN ═══════════════════╗
def main():
    print("=" * 70)
    print("  ITERATIVE PIPELINE — collect all jobs, then ATS->Sheet->ChatGPT->Excel")
    print("=" * 70)
    global SHEETS_URL
    creds = read_credentials()
    email     = creds.get("email", "")
    password  = creds.get("password", "")
    candidate = creds.get("candidate", "")
    # category may be a comma-separated LIST — iterate each one in order.
    categories = [c.strip() for c in (creds.get("category", "") or "").split(",")
                  if c.strip()]
    filter_date = creds.get("date", "")
    # Google Sheet URL — override the module default from credentials.txt.
    sheets_url_cred = (creds.get("sheets_url", "") or "").strip()
    if sheets_url_cred:
        SHEETS_URL = sheets_url_cred
    # experience: only process jobs whose YOE <= this. Blank/absent => no cap.
    max_experience = None
    exp_raw = (creds.get("experience", "") or "").strip()
    if exp_raw:
        m = re.search(r'\d+', exp_raw)
        if m:
            max_experience = int(m.group())
    print(f"[*] Builder creds: {email}")
    print(f"[*] Filters from credentials.txt -> candidate='{candidate}', "
          f"categories={categories}, date='{filter_date or '(yesterday)'}', "
          f"experience(YOE cap)={max_experience if max_experience is not None else '(none)'}")
    print(f"[*] Google Sheet URL: {SHEETS_URL[:70]}...")

    missing = [k for k in ("candidate", "category", "date") if not creds.get(k)]
    if missing:
        print(f"[!] credentials.txt is missing: {', '.join(missing)}. "
              f"Add them as 'key=value' lines and re-run.")
        return

    # Reset resume_data.xlsx to header-only only on a completely fresh run
    # (no 'done' rows in pipeline_data.xlsx). If a previous run was interrupted,
    # keep the existing resume_data.xlsx so already-written rows are preserved.
    is_fresh_run = True
    if PIPELINE_XLSX.exists():
        try:
            _wb = openpyxl.load_workbook(PIPELINE_XLSX)
            _ws = _wb.active
            for _r in range(2, _ws.max_row + 1):
                if str(_ws.cell(_r, C["status"]).value or "").strip() == "done":
                    is_fresh_run = False
                    break
        except Exception:
            pass

    if is_fresh_run or not RESUME_XLSX.exists():
        wbz = openpyxl.Workbook()
        wsz = wbz.active
        wsz.append(["Job Title", "Job Location", "Summary", "Ex Desc 1", "Ex Desc 2", "Resume Name"])
        wbz.save(RESUME_XLSX)
        print("[*] resume_data.xlsx reset to header-only (fresh run).")
    else:
        print("[*] resume_data.xlsx kept from previous run (resuming).")

    chrome_proc = None if debug_port_up() else launch_chrome()
    if chrome_proc is None:
        print("[+] Reusing running Chrome.")

    with sync_playwright() as p:
        try:
            browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{DEBUG_PORT}")
        except Exception as e:
            print(f"[~] CDP connect failed ({str(e)[:60]}); launching Chrome...")
            chrome_proc = launch_chrome()
            browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{DEBUG_PORT}")
        context = browser.contexts[0]
        for origin in ("https://firststepjob.com", "https://builder.firststepjob.com",
                       "https://docs.google.com", "https://chatgpt.com"):
            try:
                context.grant_permissions(["clipboard-read", "clipboard-write"], origin=origin)
            except Exception:
                pass
        print("[+] Clipboard permissions granted.")
        load_firststep_cookies(context)

        wb, ws = pipe_load()

        # ── STAGE 0: collect jobs for EACH category, one after another ──
        print("\n===== STAGE 0: Collect jobs for each category =====")
        apply_page = context.pages[0] if context.pages else context.new_page()
        apply_page.goto(APPLY_URL, wait_until="domcontentloaded", timeout=60000)
        try:
            apply_page.wait_for_load_state("networkidle", timeout=10000)
        except PlaywrightTimeoutError:
            pass
        for ci, cat in enumerate(categories, 1):
            print(f"\n----- Category [{ci}/{len(categories)}]: '{cat}' -----")
            # Reload the Apply Jobs page so the filter form is fresh per category.
            apply_page.goto(APPLY_URL, wait_until="domcontentloaded", timeout=60000)
            try:
                apply_page.wait_for_load_state("networkidle", timeout=10000)
            except PlaywrightTimeoutError:
                pass
            collect_all_jobs(apply_page, wb, ws, category=cat,
                             filter_date=filter_date, candidate=candidate,
                             max_experience=max_experience)

        # ── Prepare builder + sheet tabs once ────────────────────────────
        print("\n===== Preparing Builder & Sheet tabs =====")
        builder_page = context.new_page()
        builder_login_if_needed(builder_page, email, password)
        sheet_page = get_sheet_page(context)
        if sheet_page is None:
            print("[!] Google Sheet not available (sign-in needed). "
                  "Jobs will stop at the Sheet stage.")

        # ── Reset parse/response failures so they get a FRESH ChatGPT run ──
        # error:parse means we had a saved response that could not be parsed.
        # Re-parsing the same bad response would fail again, so clear it (and
        # downstream fields) to force a new ChatGPT answer on this run.
        reset_n = 0
        for r in range(2, ws.max_row + 1):
            st = str(ws.cell(r, C["status"]).value or "")
            if st in ("error:parse", "error:response"):
                pipe_set(ws, r, response=None, summary=None,
                         ex_desc_1=None, ex_desc_2=None, resume_name=None,
                         status="prompt" if ws.cell(r, C["prompt"]).value else "keywords")
                reset_n += 1
        if reset_n:
            pipe_save(wb)
            print(f"[*] Reset {reset_n} parse/response-failed job(s) for a "
                  f"fresh ChatGPT retry.")

        # ── Iterate all pending jobs ─────────────────────────────────────
        jobs = pipe_rows(ws)
        pending = [(r, v) for (r, v) in jobs if (v.get("status") or "") != "done"]
        print(f"\n===== Processing {len(pending)} pending job(s) "
              f"(of {len(jobs)} total) =====")

        done_count, fail_count = 0, 0
        for idx, (r, vals) in enumerate(pending, 1):
            jid = vals.get("job_id")
            print(f"\n----- [{idx}/{len(pending)}] Job {jid} "
                  f"({vals.get('company')}) [{vals.get('category')}] -----")
            try:
                ok = process_job(context, builder_page, sheet_page, wb, ws, r,
                                 candidate=candidate)
                if ok:
                    done_count += 1
                else:
                    fail_count += 1
            except Exception as e:
                print(f"  [!] Unhandled error on {jid}: {str(e)[:120]}")
                try:
                    pipe_set(ws, r, status=f"error:exception"); pipe_save(wb)
                except Exception:
                    pass
                fail_count += 1
                # continue to the next job

        print("\n" + "=" * 70)
        print(f"[+] RUN COMPLETE. done={done_count} failed={fail_count} "
              f"total_pending={len(pending)}")
        print(f"[+] Pipeline state: {PIPELINE_XLSX.name}")
        print(f"[+] Resumes appended to: {RESUME_XLSX.name}")

        if chrome_proc:
            chrome_proc.terminate()


if __name__ == "__main__":
    main()
