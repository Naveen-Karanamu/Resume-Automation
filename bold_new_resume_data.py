"""
Bold the ATS keywords inside the NEW resume_data.xlsx (Ex Desc 1 / Ex Desc 2
and Summary), matching each row to its job's keywords in pipeline_data.xlsx via
the company name in the Resume Name column.

Keeps ** markers as PLAIN TEXT (so resume_automation.py recreates real bold in
the builder). Idempotent, no double-bolding, whole-word matching.
"""
import re
import openpyxl
import run_all   # reuse bold_keywords + _split_keywords

COLS = ["job_id","title","company","jd","keywords","prompt","response",
        "summary","ex_desc_1","ex_desc_2","resume_name","status","category"]
C = {n: i+1 for i, n in enumerate(COLS)}

RD = "resume_data.xlsx"
PD = "pipeline_data.xlsx"
MARKER = " - Resume "


def build_company_keywords():
    """company(lower) -> keywords string (merge all keywords seen for a company)."""
    pd = openpyxl.load_workbook(PD).active
    out = {}
    for r in range(2, pd.max_row + 1):
        comp = str(pd.cell(r, C["company"]).value or "").strip().lower()
        kw = (pd.cell(r, C["keywords"]).value or "").strip()
        if not comp or not kw:
            continue
        if comp in out:
            out[comp] = out[comp] + ", " + kw   # merge keyword sets
        else:
            out[comp] = kw
    return out


def main():
    comp_kw = build_company_keywords()
    print(f"[*] Loaded keywords for {len(comp_kw)} companies from pipeline_data.")

    wb = openpyxl.load_workbook(RD)
    ws = wb.active
    bolded_rows, skipped = 0, []
    for r in range(2, ws.max_row + 1):
        rn = str(ws.cell(r, 6).value or "")
        if MARKER not in rn:
            continue
        company = rn.split(MARKER, 1)[1].strip().lower()
        kw = comp_kw.get(company)
        if not kw:
            skipped.append(company)
            continue
        changed = False
        # bold Summary (C=3), Ex Desc 1 (D=4), Ex Desc 2 (E=5)
        for col in (3, 4, 5):
            val = ws.cell(r, col).value
            if not val:
                continue
            new = run_all.bold_keywords(str(val), kw)
            if new != val:
                ws.cell(r, col, new)
                changed = True
        if changed:
            bolded_rows += 1

    wb.save(RD)
    print(f"[+] Bolded keywords in {bolded_rows} row(s). Saved {RD}.")
    if skipped:
        print(f"[~] {len(skipped)} row(s) had no matching company keywords "
              f"(left unbolded): {sorted(set(skipped))[:15]}")

    # verify
    chk = openpyxl.load_workbook(RD).active
    with_bold = sum(1 for rr in range(2, chk.max_row + 1)
                    if "**" in str(chk.cell(rr, 4).value or ""))
    print(f"[*] rows with ** in Ex Desc 1 after bolding: {with_bold}")


if __name__ == "__main__":
    main()
