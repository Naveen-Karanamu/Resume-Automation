"""
Parse chatgpt_response.txt into Summary / Ex Desc 1 / Ex Desc 2 and append
a row to resume_data.xlsx (columns C, D, E).
"""
import re
from pathlib import Path

import openpyxl

RESP = Path(__file__).parent / "chatgpt_response.txt"
XLSX = Path(__file__).parent / "resume_data.xlsx"

text = RESP.read_text(encoding="utf-8")

# Split on the major section headings the model produced.
# Headings seen:
#   "Professional Summary ... "
#   "Experience 1: ..."  (a.k.a. Responsibilities 1)
#   "Experience 2: ..."  (a.k.a. Responsibilities 2)
#   "Important ATS alignment notes"  (trailing commentary -> drop)

def section(start_pat, end_pats):
    m = re.search(start_pat, text, re.IGNORECASE)
    if not m:
        return ""
    start = m.end()
    end = len(text)
    for ep in end_pats:
        em = re.search(ep, text[start:], re.IGNORECASE)
        if em:
            end = min(end, start + em.start())
    return text[start:end].strip()

exp_heads = [r'Experience\s*1\s*[:\-]', r'Responsibilit(?:y|ies)\s*1']
summary_raw = section(r'Professional Summary[^\n]*\n', exp_heads)

resp1_raw = section(r'Experience\s*1\s*[:\-][^\n]*\n',
                    [r'Experience\s*2\s*[:\-]', r'Responsibilit(?:y|ies)\s*2'])

resp2_raw = section(r'Experience\s*2\s*[:\-][^\n]*\n',
                    [r'Important ATS', r'Before submitting', r'ATS alignment'])


def tidy_bullets(block, drop_header_lines=0):
    """
    Return the block as clean, one-per-line bullet points.
    - removes 'Writing' artifact lines
    - optionally drops a leading 'City | dates' header line
    - strips markdown hard-break trailing spaces and leading bullet glyphs
    - keeps each sentence/point on its own line (single \n between them)
    """
    lines = [ln.strip() for ln in block.splitlines()]
    lines = [ln for ln in lines if ln.lower() != "writing"]      # drop 'Writing'
    lines = [ln for ln in lines if ln != ""]                      # drop blanks

    if drop_header_lines and lines:
        removed = 0
        out = []
        for ln in lines:
            if removed < drop_header_lines and ("|" in ln and len(ln) < 80):
                removed += 1
                continue
            out.append(ln)
        lines = out

    # strip leading bullet glyphs/numbering and trailing markdown spaces
    cleaned_lines = []
    for ln in lines:
        ln = re.sub(r'^\s*[\u2022\-\*\u25cf\u25aa]\s*', '', ln)   # bullet markers
        ln = re.sub(r'^\s*\d+[\.\)]\s*', '', ln)                 # 1. 2) numbering
        ln = ln.rstrip()
        if ln:
            cleaned_lines.append(ln)
    return "\n".join(cleaned_lines).strip()


summary = tidy_bullets(summary_raw)                 # one paragraph is fine here
resp1   = tidy_bullets(resp1_raw, drop_header_lines=1)
resp2   = tidy_bullets(resp2_raw, drop_header_lines=1)

print(f"Summary : {len(summary)} chars, {summary.count(chr(10))+1} line(s)")
print(f"Resp 1  : {len(resp1)} chars, {resp1.count(chr(10))+1} line(s)")
print(f"Resp 2  : {len(resp2)} chars, {resp2.count(chr(10))+1} line(s)")
print("\n--- RESP 1 ---")
for i, ln in enumerate(resp1.split("\n"), 1):
    print(f"  {i}. {ln[:90]}")

if not (summary and resp1 and resp2):
    raise SystemExit("[!] One or more sections empty — not appending.")

from openpyxl.styles import Alignment

wb = openpyxl.load_workbook(XLSX)
ws = wb.active

# If the last row already holds this parsed content (from the earlier run),
# overwrite it in place with the properly line-separated + wrapped version
# instead of appending a duplicate.
last = ws.max_row
existing_a = ws.cell(row=last, column=1).value
existing_d = ws.cell(row=last, column=4).value
reuse_last = (
    (existing_a in (None, "")) and existing_d and
    "Architected" in str(existing_d)
)
target_row = last if reuse_last else last + 1

# A=Job Title, B=Job Location, C=Summary, D=Ex Desc 1, E=Ex Desc 2, F=Resume Name
ws.cell(row=target_row, column=3, value=summary)
ws.cell(row=target_row, column=4, value=resp1)
ws.cell(row=target_row, column=5, value=resp2)

# Enable wrap text so the embedded line breaks actually display as lines.
wrap = Alignment(wrap_text=True, vertical="top")
for col in (3, 4, 5):
    ws.cell(row=target_row, column=col).alignment = wrap

wb.save(XLSX)
action = "Overwrote" if reuse_last else "Appended"
print(f"\n[+] {action} row {target_row} in {XLSX.name} with wrap-text "
      f"and one point per line.")
