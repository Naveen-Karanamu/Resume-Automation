"""
Re-run ONLY the failed rows (status != 'done') from the sheet step.
Skips the 62 already-done jobs. For each failed row:
  switch to category tab -> paste keywords+JD -> copy prompt -> clear A9/B3 ->
  FRESH ChatGPT -> bold keywords -> parse -> record.
Then rebuild resume_data.xlsx from ALL done rows.

Requires chatgpt.com to be reachable in the debug Chrome.
"""
import time
import openpyxl
from openpyxl.styles import Alignment
import run_all as R

CAND = R.read_credentials().get("candidate", "")


def rebuild_resume_data():
    pwb = openpyxl.load_workbook("pipeline_data.xlsx"); pws = pwb.active
    out = openpyxl.Workbook(); ws = out.active; ws.title = "resume_data"
    ws.append(["Job Title", "Job Location", "Summary", "Ex Desc 1", "Ex Desc 2", "Resume Name"])
    wrap = Alignment(wrap_text=True, vertical="top")
    n = 0
    for r in range(2, pws.max_row + 1):
        summ = pws.cell(r, R.C["summary"]).value or ""
        d1 = pws.cell(r, R.C["ex_desc_1"]).value or ""
        d2 = pws.cell(r, R.C["ex_desc_2"]).value or ""
        rn = pws.cell(r, R.C["resume_name"]).value or ""
        cat = pws.cell(r, R.C["category"]).value or ""
        if not (summ and d1 and d2):
            continue
        ws.append([cat, "", summ, d1, d2, rn])
        row = ws.max_row
        for col in (3, 4, 5):
            ws.cell(row, col).alignment = wrap
        ws.row_dimensions[row].height = max(90, max(d1.count("\n")+1, d2.count("\n")+1) * 45)
        n += 1
    for cl, w in (("C", 60), ("D", 70), ("E", 70)):
        ws.column_dimensions[cl].width = w
    for a in range(10):
        try:
            out.save("resume_data.xlsx"); break
        except PermissionError:
            time.sleep(3)
    print(f"[+] Rebuilt resume_data.xlsx with {n} clean row(s).")


def main():
    from playwright.sync_api import sync_playwright
    wb, ws = R.pipe_load()

    # failed rows = not done, with keywords + jd available
    failed = [r for r in range(2, ws.max_row + 1)
              if str(ws.cell(r, R.C["status"]).value or "") != "done"
              and (ws.cell(r, R.C["keywords"]).value or "")
              and (ws.cell(r, R.C["jd"]).value or "")]
    print(f"[*] {len(failed)} failed row(s) to re-run:",
          [ws.cell(r, R.C["job_id"]).value for r in failed])
    if not failed:
        print("[+] Nothing to re-run."); rebuild_resume_data(); return

    # clear stale prompt/response on these rows so they get fresh ones
    for r in failed:
        R.pipe_set(ws, r, prompt=None, response=None, summary=None,
                   ex_desc_1=None, ex_desc_2=None, resume_name=None, status="keywords")
    R.pipe_save(wb)

    with sync_playwright() as p:
        b = p.chromium.connect_over_cdp("http://127.0.0.1:9222")
        ctx = b.contexts[0]
        for origin in ("https://docs.google.com", "https://chatgpt.com"):
            try:
                ctx.grant_permissions(["clipboard-read", "clipboard-write"], origin=origin)
            except Exception:
                pass
        sheet_page = R.get_sheet_page(ctx)
        if sheet_page is None:
            print("[!] Sheet not available."); return

        done = fail = 0
        for n, r in enumerate(failed, 1):
            jid = ws.cell(r, R.C["job_id"]).value
            cat = ws.cell(r, R.C["category"]).value or ""
            kw = ws.cell(r, R.C["keywords"]).value or ""
            jd = ws.cell(r, R.C["jd"]).value or ""
            company = ws.cell(r, R.C["company"]).value or ""
            print(f"\n----- [{n}/{len(failed)}] {jid} [{cat}] ({company}) -----")
            try:
                prompt = R.sheet_get_prompt(sheet_page, kw, jd, category=cat)
                if not prompt:
                    R.pipe_set(ws, r, status="error:prompt"); R.pipe_save(wb); fail += 1; continue
                R.pipe_set(ws, r, prompt=prompt, status="prompt"); R.pipe_save(wb)

                response = R.chatgpt_answer(ctx, prompt)
                if not response:
                    R.pipe_set(ws, r, status="error:response"); R.pipe_save(wb); fail += 1; continue
                response = R.bold_keywords(response, kw)
                R.pipe_set(ws, r, response=response, status="response"); R.pipe_save(wb)
                time.sleep(R.CHATGPT_GAP)

                summary, resp1, resp2 = R.parse_sections(response)
                if not (summary and resp1 and resp2):
                    print(f"  [!] parse incomplete (S{len(summary)} R1{bool(resp1)} R2{bool(resp2)})")
                    R.pipe_set(ws, r, summary=summary, ex_desc_1=resp1, ex_desc_2=resp2,
                               status="error:parse"); R.pipe_save(wb); fail += 1; continue
                resume_name = f"{CAND} - Resume {company}".strip() if company else f"{CAND} - Resume"
                R.pipe_set(ws, r, summary=summary, ex_desc_1=resp1, ex_desc_2=resp2,
                           resume_name=resume_name, status="done"); R.pipe_save(wb)
                print(f"  [+] DONE {jid} -> {resume_name} (R1 {resp1.count(chr(10))+1}, R2 {resp2.count(chr(10))+1} lines)")
                done += 1
            except Exception as e:
                print(f"  [!] error on {jid}: {str(e)[:110]}")
                try:
                    R.pipe_set(ws, r, status="error:exception"); R.pipe_save(wb)
                except Exception:
                    pass
                fail += 1

        print(f"\n[+] Re-run complete. done={done} failed={fail}")

    rebuild_resume_data()


if __name__ == "__main__":
    main()
