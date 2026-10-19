"""Rebuild resume_data.xlsx from the (bolded) responses in pipeline_data.xlsx.
Keeps ** markers as plain text so resume_automation.py can recreate real bold.
Job Title (col A) = category (fallback to scraped title)."""
import openpyxl
from openpyxl.styles import Alignment
import run_all   # reuse parse_sections

COLS = ["job_id","title","company","jd","keywords","prompt","response",
        "summary","ex_desc_1","ex_desc_2","resume_name","status","category"]
C = {n: i+1 for i, n in enumerate(COLS)}

pws = openpyxl.load_workbook("pipeline_data.xlsx").active

out = openpyxl.Workbook(); ws = out.active; ws.title = "resume_data"
ws.append(["Job Title","Job Location","Summary","Ex Desc 1","Ex Desc 2","Resume Name"])
wrap = Alignment(wrap_text=True, vertical="top")

count = 0
for r in range(2, pws.max_row + 1):
    if pws.cell(r, C["status"]).value != "done":
        continue
    resp = pws.cell(r, C["response"]).value or ""
    if not resp:
        continue
    summary, d1, d2 = run_all.parse_sections(resp)
    if not (summary and d1 and d2):
        continue
    rn = pws.cell(r, C["resume_name"]).value or ""
    job_title = pws.cell(r, C["category"]).value or pws.cell(r, C["title"]).value or ""
    ws.append([job_title, "", summary, d1, d2, rn])
    row = ws.max_row
    for col in (3,4,5):
        ws.cell(row, col).alignment = wrap
    ml = max(d1.count("\n")+1, d2.count("\n")+1)
    ws.row_dimensions[row].height = max(90, ml*45)
    count += 1

for cl,w in (("C",60),("D",70),("E",70)):
    ws.column_dimensions[cl].width = w
out.save("resume_data.xlsx")
print(f"Rebuilt resume_data.xlsx: {count} rows.")
# verify ** survived
chk = openpyxl.load_workbook("resume_data.xlsx").active
withbold = sum(1 for rr in range(2, chk.max_row+1) if "**" in str(chk.cell(rr,4).value or ""))
print(f"rows with ** in Ex Desc 1: {withbold}")
