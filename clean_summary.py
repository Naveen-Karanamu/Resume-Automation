"""Strip a leading 'ChatGPT said:' prefix from the Summary (C), Ex Desc 1 (D),
and Ex Desc 2 (E) columns of resume_data.xlsx."""
import re
import openpyxl

RD = "resume_data.xlsx"
PREFIX = re.compile(r'^\s*chatgpt said\s*:?\s*', re.IGNORECASE)

wb = openpyxl.load_workbook(RD)
ws = wb.active
fixed = 0
for r in range(2, ws.max_row + 1):
    for col in (3, 4, 5):
        v = ws.cell(r, col).value
        if v and PREFIX.match(str(v)):
            ws.cell(r, col, PREFIX.sub('', str(v)))
            fixed += 1
wb.save(RD)
print(f"[+] Stripped 'ChatGPT said:' from {fixed} cell(s). Saved {RD}.")

# verify none remain
chk = openpyxl.load_workbook(RD).active
left = sum(1 for r in range(2, chk.max_row + 1) for c in (3,4,5)
           if str(chk.cell(r,c).value or '').lower().startswith('chatgpt said'))
print(f"[*] cells still starting with 'ChatGPT said': {left}")
